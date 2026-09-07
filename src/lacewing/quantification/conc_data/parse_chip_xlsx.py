"""Parse per-chip LC96 xlsx exports (Chip 1.xlsx, Chip 2.xlsx) into a
per-concentration plate-TTP dict.

Per supervisor guidance (2026-08-14):
  - There is now ONE xlsx per chip.
  - For each chip, the plate TTP for a concentration is the MEAN of all Cq
    values in rows whose `Position` matches that Kp_N label (Kp_3 / Kp_4 /
    Kp_5 / Kp_6). NTC rows are excluded even if they carry a stray Cq.
  - Both chip wells at the same concentration inherit the SAME plate TTP
    (technically identical replicates on the chip).

Output shape:
    {
      "chip1": {
        "source": "...",
        "cycle_time_min": 0.5,
        "per_conc_kp": {
          "Kp_3": {"cq_mean": ..., "ttp_min": ..., "n_positive": ..., "raw_cqs": [...]},
          "Kp_4": {...},
          ...
        },
      },
      "chip2": {...},
    }
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
from lacewing._paths import LACEWING_PKG_DIR, DATA_ROOT


PROJECT_ROOT = DATA_ROOT  # was: parents[3]
DATA_ROOT = PROJECT_ROOT / "Data" / "Concentration Data Experiment"
CHIP_XLSX = {
    "chip1": DATA_ROOT / "Chip 1.xlsx",
    "chip2": DATA_ROOT / "Chip 2.xlsx",
}
OUT_JSON = DATA_ROOT / "plate_features.json"

CYCLE_TIME_MIN = 0.5   # 30-sec LC96 LAMP cycles

KP_LABELS = ("Kp_3", "Kp_4", "Kp_5", "Kp_6")


def _f(x) -> float | None:
    if x is None:
        return None
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip()
    if not s or s == "-":
        return None
    try:
        return float(s)
    except ValueError:
        return None


def parse_one_chip(path: Path) -> dict:
    if not path.exists():
        raise FileNotFoundError(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Sheet1"]
    header = [str(c).strip() if c is not None else "" for c in next(ws.iter_rows(values_only=True))]

    def col(name: str) -> int:
        for i, h in enumerate(header):
            if h == name:
                return i
        raise KeyError(f"column '{name}' not found in {path.name}. Have: {header}")

    ci_pos    = col("Position")
    ci_cq     = col("Cq")
    ci_call   = col("Call")
    ci_sample = col("Sample Name")
    ci_gene   = col("Gene Name")

    per_kp: dict[str, dict] = {k: {"raw_cqs": [], "raw_rows": []} for k in KP_LABELS}
    ntc_rows: list[dict] = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None or all(c is None for c in row):
            continue
        pos    = (str(row[ci_pos]).strip()   if row[ci_pos]    is not None else "")
        cq     = _f(row[ci_cq])
        call   = (str(row[ci_call]).strip()  if row[ci_call]   is not None else "")
        sample = (str(row[ci_sample]).strip() if row[ci_sample] is not None else "")
        gene   = (str(row[ci_gene]).strip()   if row[ci_gene]   is not None else "")

        rec = {"cq": cq, "call": call, "sample_name": sample, "gene_name": gene}

        if pos in per_kp:
            per_kp[pos]["raw_rows"].append(rec)
            if cq is not None:
                per_kp[pos]["raw_cqs"].append(cq)
        elif pos == "NTC":
            ntc_rows.append({"cq": cq, "call": call, "sample_name": sample, "gene_name": gene})

    # Aggregate per Kp_N.
    for kp, entry in per_kp.items():
        cqs = entry["raw_cqs"]
        cq_mean = (sum(cqs) / len(cqs)) if cqs else None
        entry["n_replicates"] = len(entry["raw_rows"])
        entry["n_positive"]   = len(cqs)
        entry["cq_mean"]      = cq_mean
        entry["ttp_min"]      = (cq_mean * CYCLE_TIME_MIN) if cq_mean is not None else None

    return {
        "source": str(path.relative_to(PROJECT_ROOT)),
        "cycle_time_min": CYCLE_TIME_MIN,
        "notes": (
            "Cq mean is computed across ONLY the Kp_N-labelled rows for each "
            "concentration (NTC rows excluded even if they carry a stray Cq). "
            "Both chip wells at the same concentration get this same TTP."
        ),
        "per_conc_kp": per_kp,
        "ntc_rows_ignored": ntc_rows,
    }


def parse_all() -> dict:
    return {chip_key: parse_one_chip(path) for chip_key, path in CHIP_XLSX.items()}


def main() -> None:
    payload = parse_all()
    OUT_JSON.write_text(json.dumps(payload, indent=2))
    print(f"[ok] wrote {OUT_JSON}")
    for chip_key, chip in payload.items():
        print(f"\n=== {chip_key} ({chip['source']}) ===")
        for kp in KP_LABELS:
            e = chip["per_conc_kp"][kp]
            cq_str  = "None" if e["cq_mean"] is None else f"{e['cq_mean']:.3f}"
            ttp_str = "None" if e["ttp_min"] is None else f"{e['ttp_min']:.3f} min"
            print(f"  {kp}: n_pos={e['n_positive']}/{e['n_replicates']}  "
                  f"Cq_mean={cq_str}  TTP={ttp_str}  raw={e['raw_cqs']}")
        if chip["ntc_rows_ignored"]:
            print(f"  NTC rows ignored (n={len(chip['ntc_rows_ignored'])}): "
                  f"{[r['cq'] for r in chip['ntc_rows_ignored']]}")


if __name__ == "__main__":
    main()
